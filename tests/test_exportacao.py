from io import BytesIO

import pytest
from docx import Document as abrir_docx

from docagent.models import Bloco, Documento, Organizacao
from docagent.renderers import exportar, para_markdown
from docagent.renderers.tema import agrupar_assinaturas, larguras_colunas, segmentos


def _texto_docx(conteudo: bytes) -> str:
    d = abrir_docx(BytesIO(conteudo))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for linha in t.rows:
            partes += [c.text for c in linha.cells]
    return "\n".join(partes)


@pytest.mark.parametrize("formato", ["md", "docx", "pdf"])
@pytest.mark.parametrize("com_org", [True, False])
def test_exporta_todos_os_formatos(documento, org, formato, com_org):
    dados = exportar(documento, org if com_org else None, formato)
    assert len(dados) > 500
    if formato == "pdf":
        assert dados.startswith(b"%PDF")
    if formato == "docx":
        assert documento.titulo in _texto_docx(dados)
    if formato == "md":
        assert dados.decode("utf-8").count(f"# {documento.titulo}") == 1


def test_docx_tem_cabecalho_da_organizacao(documento, org):
    d = abrir_docx(BytesIO(exportar(documento, org, "docx")))
    cabecalho = d.sections[0].header
    textos = [c.text for t in cabecalho.tables for linha in t.rows for c in linha.cells]
    assert any("Empresa Exemplo Ltda." in t for t in textos)


def test_markdown_formata_blocos():
    doc = Documento(
        titulo="Teste",
        subtitulo="",
        metadados=[],
        pendencias=[],
        blocos=[
            Bloco(tipo="tabela", texto="", itens=[], colunas=["A", "B"], linhas=[["1|x", "2"], ["3"]]),
            Bloco(tipo="lista_numerada", texto="", itens=["um", "dois"], colunas=[], linhas=[]),
            Bloco(tipo="destaque", texto="Atenção", itens=[], colunas=[], linhas=[]),
        ],
    )
    md = para_markdown(doc, Organizacao())
    assert "| A | B |" in md
    assert "| 1\\|x | 2 |" in md
    assert "| 3 |  |" in md  # linha incompleta é completada
    assert "2. dois" in md
    assert "> Atenção" in md


def test_pdf_aceita_caracteres_especiais_e_tabela_vazia(org):
    doc = Documento(
        titulo="Símbolos → ≥ ✓ & <tags>",
        subtitulo="",
        metadados=[],
        pendencias=[],
        blocos=[
            Bloco(tipo="paragrafo", texto="R$ 1.000 & **50%** <b>", itens=[], colunas=[], linhas=[]),
            Bloco(tipo="tabela", texto="", itens=[], colunas=[], linhas=[]),
        ],
    )
    assert exportar(doc, org, "pdf").startswith(b"%PDF")
    assert _texto_docx(exportar(doc, org, "docx"))


def test_logo_invalido_nao_quebra_exportacao(documento):
    org = Organizacao(nome="X", logo=b"nao e imagem")
    assert exportar(documento, org, "pdf").startswith(b"%PDF")
    assert exportar(documento, org, "docx")


def test_formato_desconhecido(documento):
    with pytest.raises(ValueError):
        exportar(documento, None, "odt")  # type: ignore[arg-type]


def test_segmentos_negrito():
    assert segmentos("a **b** c") == [("a ", False), ("b", True), (" c", False)]
    assert segmentos("sem negrito") == [("sem negrito", False)]


def test_larguras_respeitam_total_e_minimos():
    colunas = ["Objetivo", "Ação", "Tipo"]
    linhas = [["1", "Uma descrição bastante longa de uma ação de desenvolvimento", "70"]]
    larguras = larguras_colunas(colunas, linhas, 100.0, lambda t, b: len(t) * 1.0, padding=2)
    assert sum(larguras) == pytest.approx(100.0)
    assert larguras[0] >= len("Objetivo") + 2
    assert larguras[1] > larguras[2]


def test_agrupar_assinaturas():
    def b(tipo):
        return Bloco(tipo=tipo, texto="", itens=[], colunas=[], linhas=[])

    grupos = list(agrupar_assinaturas([b("paragrafo"), b("assinatura"), b("assinatura"), b("lista")]))
    assert [type(g).__name__ for g in grupos] == ["Bloco", "list", "Bloco"]
    assert len(grupos[1]) == 2


def test_moeda_nao_quebra_em_tabelas():
    from docagent.renderers.tema import nao_quebrar, normalizar_tabela

    assert nao_quebrar("R$ 18.300") == "R$\u00a018.300"
    assert nao_quebrar("+R$ 2.150 e R$ −7") == "+R$\u00a02.150 e R$\u00a0−7"
    assert normalizar_tabela(["A"], [["R$ 10"]]) == [["R$\u00a010"]]
