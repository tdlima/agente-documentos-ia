"""Exportação para Word (.docx) com python-docx."""

from __future__ import annotations

from io import BytesIO

from docx import Document as novo_docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..models import Bloco, Documento, Organizacao
from .tema import (
    CINZA,
    FUNDO_DESTAQUE,
    LATAO,
    MARINHO,
    TEXTO,
    agrupar_assinaturas,
    larguras_colunas,
    normalizar_tabela,
    segmentos,
)

LARGURA_UTIL_CM = 16.0  # A4 (21 cm) menos margens de 2,5 cm

FONTE = "Calibri"


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# Ordem dos elementos exigida pelo schema do Word (OOXML) dentro de <w:pPr>.
_APOS_PBDR = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
_APOS_SHD_CELULA = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _novo_shd(cor: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    return shd


def _sombrear_celula(celula, cor: str) -> None:
    celula._tc.get_or_add_tcPr().insert_element_before(_novo_shd(cor), *_APOS_SHD_CELULA)


def _sombrear_paragrafo(paragrafo, cor: str) -> None:
    paragrafo._p.get_or_add_pPr().insert_element_before(_novo_shd(cor), *_APOS_PBDR[1:])


def _borda_paragrafo(paragrafo, lado: str, cor: str, espessura: int = 12) -> None:
    pPr = paragrafo._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(bdr, *_APOS_PBDR)
    borda = OxmlElement(f"w:{lado}")
    borda.set(qn("w:val"), "single")
    borda.set(qn("w:sz"), str(espessura))
    borda.set(qn("w:space"), "6")
    borda.set(qn("w:color"), cor)
    bdr.append(borda)


def _campo(paragrafo, instrucao: str) -> None:
    """Insere um campo do Word (ex.: PAGE, NUMPAGES)."""
    run = paragrafo.add_run()
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(CINZA)
    for tipo, texto in (("begin", None), (None, instrucao), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def _escrever(paragrafo, texto: str, tamanho: float | None = None, cor: str | None = None) -> None:
    for trecho, negrito in segmentos(texto):
        run = paragrafo.add_run(trecho)
        run.bold = negrito or None
        if tamanho:
            run.font.size = Pt(tamanho)
        if cor:
            run.font.color.rgb = _rgb(cor)


def _configurar_estilos(d) -> None:
    normal = d.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(TEXTO)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nome, tamanho in (("Title", 20), ("Heading 2", 13)):
        estilo = d.styles[nome]
        estilo.font.name = FONTE
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = _rgb(MARINHO)
        fontes = estilo.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fontes.attrib.pop(qn(attr), None)  # usa a fonte definida, não a do tema
    d.styles["Heading 2"].paragraph_format.space_before = Pt(14)
    d.styles["Heading 2"].paragraph_format.space_after = Pt(4)
    d.styles["Heading 2"].paragraph_format.keep_with_next = True


def _definir_larguras(tabela, larguras_cm: list[float]) -> None:
    """Fixa larguras das colunas (células + tblGrid), respeitado por Word e LibreOffice."""
    tabela.autofit = False
    grade = tabela._tbl.tblGrid
    for col, largura in zip(grade.findall(qn("w:gridCol")), larguras_cm):
        col.set(qn("w:w"), str(int(Cm(largura).twips)))
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras_cm):
            celula.width = Cm(largura)


def _cabecalho_rodape(d, org: Organizacao) -> None:
    secao = d.sections[0]
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    secao.top_margin = Cm(2.5)
    secao.bottom_margin = Cm(2)
    secao.left_margin = secao.right_margin = Cm((21 - LARGURA_UTIL_CM) / 2)
    secao.header_distance = Cm(1)

    if org.nome or org.logo:
        cab = secao.header
        # tabela sem bordas: [logo | nome + contatos]
        t = cab.add_table(rows=1, cols=2, width=Cm(LARGURA_UTIL_CM))
        cel_logo, cel_texto = t.rows[0].cells
        largura_logo = 2.0 if org.logo else 0.1
        _definir_larguras(t, [largura_logo, LARGURA_UTIL_CM - largura_logo])
        if org.logo:
            try:
                cel_logo.paragraphs[0].add_run().add_picture(BytesIO(org.logo), height=Cm(1.5))
            except Exception:  # imagem inválida: segue sem logo
                pass
        p = cel_texto.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(org.nome)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = _rgb(MARINHO)
        for linha in org.linhas_cabecalho():
            pl = cel_texto.add_paragraph()
            pl.paragraph_format.space_after = Pt(0)
            pl.paragraph_format.line_spacing = 1.0
            r = pl.add_run(linha)
            r.font.size = Pt(8)
            r.font.color.rgb = _rgb(CINZA)
        # o parágrafo padrão do cabeçalho fica abaixo da tabela, com a linha dourada
        divisor = cab.paragraphs[0]
        cab._element.remove(divisor._p)
        cab._element.append(divisor._p)
        divisor.paragraph_format.space_after = Pt(0)
        _borda_paragrafo(divisor, "bottom", LATAO, 8)

    rod = secao.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for parte in ("Página ", "PAGE", " de ", "NUMPAGES"):
        if parte in ("PAGE", "NUMPAGES"):
            _campo(rod, parte)
        else:
            run = rod.add_run(parte)
            run.font.size = Pt(8)
            run.font.color.rgb = _rgb(CINZA)


def _tabela(d, bloco: Bloco) -> None:
    colunas = bloco.colunas
    linhas = normalizar_tabela(colunas, bloco.linhas)
    n = len(colunas) or (len(linhas[0]) if linhas else 0)
    if n == 0:
        return
    t = d.add_table(rows=1 if colunas else 0, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if colunas:
        for celula, texto in zip(t.rows[0].cells, colunas):
            celula.text = ""
            _escrever(celula.paragraphs[0], texto, 10, "FFFFFF")
            for run in celula.paragraphs[0].runs:
                run.bold = True
            _sombrear_celula(celula, MARINHO)
        # repete o cabeçalho em quebras de página
        trPr = t.rows[0]._tr.get_or_add_trPr()
        cabecalho = OxmlElement("w:tblHeader")
        cabecalho.set(qn("w:val"), "true")
        trPr.append(cabecalho)
    for i, linha in enumerate(linhas):
        celulas = t.add_row().cells
        for celula, texto in zip(celulas, linha):
            celula.text = ""
            _escrever(celula.paragraphs[0], texto, 10)
            if i % 2:
                _sombrear_celula(celula, "F3F4F6")

    # estimativa para Calibri 10 pt: ~0,19 cm por caractere (0,21 em negrito)
    def medir(texto: str, negrito: bool) -> float:
        return len(texto) * (0.21 if negrito else 0.19)

    _definir_larguras(t, larguras_colunas(colunas, linhas, LARGURA_UTIL_CM, medir, padding=0.45))
    d.add_paragraph()


def _assinaturas(d, blocos: list[Bloco]) -> None:
    """Assinaturas lado a lado (duas por linha) em uma tabela sem bordas."""
    colunas = 1 if len(blocos) == 1 else 2
    espaco = d.add_paragraph()
    espaco.paragraph_format.space_before = Pt(24)
    espaco.paragraph_format.keep_with_next = True
    for i in range(0, len(blocos), colunas):
        t = d.add_table(rows=1, cols=colunas)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _definir_larguras(t, [LARGURA_UTIL_CM / colunas] * colunas)
        for celula, b in zip(t.rows[0].cells, blocos[i : i + colunas]):
            linha = celula.paragraphs[0]
            linha.alignment = WD_ALIGN_PARAGRAPH.CENTER
            linha.paragraph_format.space_before = Pt(30)
            linha.paragraph_format.keep_with_next = True
            linha.add_run("_" * 34)
            nome = celula.add_paragraph()
            nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
            nome.paragraph_format.space_after = Pt(0)
            nome.add_run(b.texto).bold = True
            if b.itens and b.itens[0]:
                cargo = celula.add_paragraph()
                cargo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cargo.add_run(b.itens[0])
                r.font.color.rgb = _rgb(CINZA)
        # impede que a linha da tabela seja dividida entre páginas
        trPr = t.rows[0]._tr.get_or_add_trPr()
        nao_dividir = OxmlElement("w:cantSplit")
        trPr.append(nao_dividir)


def para_docx(doc: Documento, org: Organizacao | None = None) -> bytes:
    org = org or Organizacao()
    d = novo_docx()
    _configurar_estilos(d)
    _cabecalho_rodape(d, org)
    d.core_properties.title = doc.titulo
    d.core_properties.author = org.nome or "DocAgent"

    titulo = d.add_paragraph(style="Title")
    titulo.add_run(doc.titulo)
    _borda_paragrafo(titulo, "bottom", LATAO, 12)
    if doc.subtitulo:
        sub = d.add_paragraph()
        r = sub.add_run(doc.subtitulo)
        r.italic = True
        r.font.color.rgb = _rgb(CINZA)

    for m in doc.metadados:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{m.rotulo}: ")
        r.bold = True
        r.font.color.rgb = _rgb(MARINHO)
        _escrever(p, m.valor)
    if doc.metadados:
        d.add_paragraph()

    for b in agrupar_assinaturas(doc.blocos):
        if isinstance(b, list):
            _assinaturas(d, b)
        elif b.tipo == "titulo":
            d.add_heading(b.texto, level=2)
        elif b.tipo == "paragrafo":
            for parte in [x for x in b.texto.split("\n") if x.strip()] or [""]:
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _escrever(p, parte)
        elif b.tipo in ("lista", "lista_numerada"):
            for i, item in enumerate(b.itens, 1):
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.space_after = Pt(2)
                marcador = f"{i}.\t" if b.tipo == "lista_numerada" else "•\t"
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
                r = p.add_run(marcador)
                r.font.color.rgb = _rgb(LATAO)
                r.bold = True
                _escrever(p, item)
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif b.tipo == "tabela":
            _tabela(d, b)
        elif b.tipo == "destaque":
            p = d.add_paragraph()
            _borda_paragrafo(p, "left", LATAO, 24)
            _sombrear_paragrafo(p, FUNDO_DESTAQUE)
            p.paragraph_format.left_indent = Cm(0.3)
            _escrever(p, b.texto)

    # remove parágrafo vazio final para evitar página em branco
    if d.paragraphs and not d.paragraphs[-1].text.strip() and len(d.paragraphs) > 1:
        ultimo = d.paragraphs[-1]._p
        if not ultimo.xpath(".//w:br[@w:type='page']"):
            ultimo.getparent().remove(ultimo)

    buffer = BytesIO()
    d.save(buffer)
    return buffer.getvalue()
